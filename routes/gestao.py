from flask import Blueprint, render_template, request, redirect, url_for, flash, session, send_file, jsonify
from models import db, CicloRevisao, IndicadorDado
from data import INDICADORES
from datetime import datetime
import uuid
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

gestao_bp = Blueprint('gestao', __name__)


def verificar_autenticacao():
    if 'autenticado' not in session:
        flash('Acesso negado. Faça login primeiro.', 'danger')
        return False
    return True


@gestao_bp.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        from app import app
        password = request.form.get('password', '')
        if password == app.config['GESTAO_PASSWORD']:
            session['autenticado'] = True
            flash('Login realizado com sucesso!', 'success')
            return redirect(url_for('gestao.index'))
        else:
            flash('Password incorreta.', 'danger')

    return render_template('gestao/login.html')


@gestao_bp.route('/logout', methods=['GET'])
def logout():
    session.clear()
    flash('Logout realizado.', 'success')
    return redirect(url_for('gestao.login'))


@gestao_bp.route('/', methods=['GET'])
def index():
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    estado = request.args.get('estado', 'todos')
    page = request.args.get('page', 1, type=int)

    query = CicloRevisao.query.order_by(CicloRevisao.criado_em.desc())

    if estado == 'pendente':
        query = query.filter_by(estado='pendente')
    elif estado == 'revisto':
        query = query.filter_by(estado='revisto')

    ciclos = query.paginate(page=page, per_page=10)

    return render_template('gestao/index.html', ciclos=ciclos, estado=estado)


@gestao_bp.route('/ciclo/novo', methods=['GET', 'POST'])
def ciclo_novo():
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    if request.method == 'POST':
        titulo = request.form.get('titulo', '').strip()

        if not titulo:
            flash('Título do ciclo é obrigatório.', 'danger')
            return redirect(url_for('gestao.ciclo_novo'))

        ciclo = CicloRevisao(
            titulo=titulo,
            token=str(uuid.uuid4()),
            estado='pendente'
        )
        db.session.add(ciclo)
        db.session.flush()

        for grupo in INDICADORES:
            for ind_data in grupo['indicadores']:
                ind = IndicadorDado(
                    ciclo_id=ciclo.id,
                    codigo=ind_data['codigo'],
                    grupo=grupo['grupo_id'],
                    tipo=ind_data['tipo'],
                    nome_original=ind_data['nome'],
                    definicao=ind_data['definicao'],
                    meta_original=ind_data['meta'],
                )
                db.session.add(ind)

        db.session.commit()
        flash(f'Ciclo "{titulo}" criado com sucesso!', 'success')
        return redirect(url_for('gestao.ciclo_detalhe', id=ciclo.id))

    return render_template('gestao/ciclo_novo.html', grupos=INDICADORES)


@gestao_bp.route('/ciclo/<int:id>', methods=['GET'])
def ciclo_detalhe(id):
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    ciclo = CicloRevisao.query.get_or_404(id)

    indicadores_por_grupo = {}
    for ind in ciclo.indicadores:
        if ind.grupo not in indicadores_por_grupo:
            indicadores_por_grupo[ind.grupo] = []
        indicadores_por_grupo[ind.grupo].append(ind)

    return render_template('gestao/ciclo.html', ciclo=ciclo, indicadores_por_grupo=indicadores_por_grupo)


@gestao_bp.route('/ciclo/<int:id>/editar', methods=['GET', 'POST'])
def ciclo_editar(id):
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    ciclo = CicloRevisao.query.get_or_404(id)

    if ciclo.estado == 'revisto':
        flash('Não pode editar um ciclo já revisto.', 'warning')
        return redirect(url_for('gestao.ciclo_detalhe', id=id))

    if request.method == 'POST':
        for ind in ciclo.indicadores:
            resultado_actual = request.form.get(f'resultado_actual_{ind.id}', '').strip()
            projectos = request.form.get(f'projectos_{ind.id}', '').strip()
            notas_equipa = request.form.get(f'notas_equipa_{ind.id}', '').strip()

            ind.resultado_actual = resultado_actual if resultado_actual else None
            ind.projectos = projectos if projectos else None
            ind.notas_equipa = notas_equipa if notas_equipa else None

        db.session.commit()
        flash('Dados actualizados com sucesso!', 'success')
        return redirect(url_for('gestao.ciclo_detalhe', id=id))

    indicadores_por_grupo = {}
    for ind in ciclo.indicadores:
        if ind.grupo not in indicadores_por_grupo:
            indicadores_por_grupo[ind.grupo] = []
        indicadores_por_grupo[ind.grupo].append(ind)

    return render_template('gestao/ciclo_editar.html', ciclo=ciclo, indicadores_por_grupo=indicadores_por_grupo, grupos=INDICADORES)


@gestao_bp.route('/ciclo/<int:id>/apagar', methods=['POST'])
def ciclo_apagar(id):
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    ciclo = CicloRevisao.query.get_or_404(id)
    titulo = ciclo.titulo
    db.session.delete(ciclo)
    db.session.commit()
    flash(f'Ciclo "{titulo}" apagado com sucesso.', 'success')
    return redirect(url_for('gestao.index'))


@gestao_bp.route('/ciclo/<int:id>/link', methods=['GET'])
def ciclo_link(id):
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    ciclo = CicloRevisao.query.get_or_404(id)
    from flask import request as flask_request
    link = flask_request.url_root.rstrip('/') + url_for('revisao.form', token=ciclo.token)

    return render_template('gestao/ciclo_link.html', ciclo=ciclo, link=link)


def criar_excel_ciclo(ciclo):
    wb = Workbook()
    ws_resumo = wb.active
    ws_resumo.title = 'Resumo'

    header_fill = PatternFill(start_color='1a5276', end_color='1a5276', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    ws_resumo['A1'] = 'CICLO DE REVISÃO'
    ws_resumo['A1'].font = Font(bold=True, size=14)
    ws_resumo['A2'] = 'Título:'
    ws_resumo['B2'] = ciclo.titulo
    ws_resumo['A3'] = 'Data de Criação:'
    ws_resumo['B3'] = ciclo.criado_em.strftime('%d/%m/%Y %H:%M')
    ws_resumo['A4'] = 'Estado:'
    ws_resumo['B4'] = 'Revisto' if ciclo.estado == 'revisto' else 'Pendente'

    if ciclo.estado == 'revisto':
        ws_resumo['A5'] = 'Director:'
        ws_resumo['B5'] = ciclo.nome_director
        ws_resumo['A6'] = 'Data de Revisão:'
        ws_resumo['B6'] = ciclo.data_revisao.strftime('%d/%m/%Y %H:%M')
        ws_resumo['A7'] = 'Observações Gerais:'
        ws_resumo['B7'] = ciclo.observacoes_director or '—'

    confirmados = sum(1 for ind in ciclo.indicadores if ind.confirmado)
    editados = sum(1 for ind in ciclo.indicadores if ind.foi_editado_pelo_director)
    eliminacao = sum(1 for ind in ciclo.indicadores if ind.propor_eliminacao)

    ws_resumo['A9'] = 'RESUMO'
    ws_resumo['A9'].font = Font(bold=True, size=12)
    ws_resumo['A10'] = 'Total de Indicadores:'
    ws_resumo['B10'] = len(ciclo.indicadores)
    ws_resumo['A11'] = 'Confirmados:'
    ws_resumo['B11'] = confirmados
    ws_resumo['A12'] = 'Editados pelo Director:'
    ws_resumo['B12'] = editados
    ws_resumo['A13'] = 'Propostos para Eliminação:'
    ws_resumo['B13'] = eliminacao

    ws_resumo.column_dimensions['A'].width = 30
    ws_resumo.column_dimensions['B'].width = 40

    ws_ind = wb.create_sheet('Indicadores')
    headers = [
        'Grupo', 'Código', 'Tipo', 'Nome Original', 'Nome Revisto',
        'Meta Original', 'Meta Revista', 'Resultado Actual', 'Projectos',
        'Notas Equipa', 'Notas Director', 'Confirmado', 'Propor Eliminação', 'Justificação'
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws_ind.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    for row_num, ind in enumerate(ciclo.indicadores, 2):
        row_data = [
            ind.grupo,
            ind.codigo,
            'Monitoria' if ind.tipo == 'M' else 'Avaliação',
            ind.nome_original,
            ind.nome_editado or '—',
            ind.meta_original,
            ind.meta_editada or '—',
            ind.resultado_actual or '—',
            ind.projectos or '—',
            ind.notas_equipa or '—',
            ind.notas_director or '—',
            'Sim' if ind.confirmado else 'Não',
            'Sim' if ind.propor_eliminacao else 'Não',
            ind.justificacao_eliminacao or '—'
        ]

        for col_num, value in enumerate(row_data, 1):
            cell = ws_ind.cell(row=row_num, column=col_num)
            cell.value = value
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical='top')

            if ind.propor_eliminacao:
                cell.fill = PatternFill(start_color='fadbd8', end_color='fadbd8', fill_type='solid')
            elif ind.foi_editado_pelo_director:
                cell.fill = PatternFill(start_color='fef9e7', end_color='fef9e7', fill_type='solid')
            elif ind.confirmado:
                cell.fill = PatternFill(start_color='eafaf1', end_color='eafaf1', fill_type='solid')

    for col_num in range(1, len(headers) + 1):
        ws_ind.column_dimensions[get_column_letter(col_num)].width = 20

    return wb


@gestao_bp.route('/ciclo/<int:id>/exportar', methods=['GET'])
def ciclo_exportar(id):
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    ciclo = CicloRevisao.query.get_or_404(id)
    wb = criar_excel_ciclo(ciclo)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'mci_indicadores_{ciclo.id}.xlsx'
    )


@gestao_bp.route('/exportar-todos', methods=['GET'])
def exportar_todos():
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    ciclos = CicloRevisao.query.all()
    wb = Workbook()
    wb.remove(wb.active)

    for ciclo in ciclos:
        titulo_truncado = ciclo.titulo[:31]
        ws = wb.create_sheet(titulo_truncado)

        header_fill = PatternFill(start_color='1a5276', end_color='1a5276', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        headers = [
            'Ciclo', 'Director', 'Grupo', 'Código', 'Tipo', 'Nome Original', 'Nome Revisto',
            'Meta Original', 'Meta Revista', 'Resultado Actual', 'Projectos',
            'Notas Equipa', 'Notas Director', 'Confirmado', 'Propor Eliminação', 'Justificação'
        ]

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border

        for row_num, ind in enumerate(ciclo.indicadores, 2):
            row_data = [
                ciclo.titulo,
                ciclo.nome_director or '—',
                ind.grupo,
                ind.codigo,
                'Monitoria' if ind.tipo == 'M' else 'Avaliação',
                ind.nome_original,
                ind.nome_editado or '—',
                ind.meta_original,
                ind.meta_editada or '—',
                ind.resultado_actual or '—',
                ind.projectos or '—',
                ind.notas_equipa or '—',
                ind.notas_director or '—',
                'Sim' if ind.confirmado else 'Não',
                'Sim' if ind.propor_eliminacao else 'Não',
                ind.justificacao_eliminacao or '—'
            ]

            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical='top')

                if ind.propor_eliminacao:
                    cell.fill = PatternFill(start_color='fadbd8', end_color='fadbd8', fill_type='solid')
                elif ind.foi_editado_pelo_director:
                    cell.fill = PatternFill(start_color='fef9e7', end_color='fef9e7', fill_type='solid')
                elif ind.confirmado:
                    cell.fill = PatternFill(start_color='eafaf1', end_color='eafaf1', fill_type='solid')

        for col_num in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_num)].width = 18

    ws_consolidado = wb.create_sheet('Consolidado')
    headers = [
        'Ciclo', 'Director', 'Grupo', 'Código', 'Tipo', 'Nome Original', 'Nome Revisto',
        'Meta Original', 'Meta Revista', 'Resultado Actual', 'Projectos',
        'Notas Equipa', 'Notas Director', 'Confirmado', 'Propor Eliminação', 'Justificação'
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws_consolidado.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    row_num = 2
    for ciclo in ciclos:
        for ind in ciclo.indicadores:
            row_data = [
                ciclo.titulo,
                ciclo.nome_director or '—',
                ind.grupo,
                ind.codigo,
                'Monitoria' if ind.tipo == 'M' else 'Avaliação',
                ind.nome_original,
                ind.nome_editado or '—',
                ind.meta_original,
                ind.meta_editada or '—',
                ind.resultado_actual or '—',
                ind.projectos or '—',
                ind.notas_equipa or '—',
                ind.notas_director or '—',
                'Sim' if ind.confirmado else 'Não',
                'Sim' if ind.propor_eliminacao else 'Não',
                ind.justificacao_eliminacao or '—'
            ]

            for col_num, value in enumerate(row_data, 1):
                cell = ws_consolidado.cell(row=row_num, column=col_num)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical='top')

                if ind.propor_eliminacao:
                    cell.fill = PatternFill(start_color='fadbd8', end_color='fadbd8', fill_type='solid')
                elif ind.foi_editado_pelo_director:
                    cell.fill = PatternFill(start_color='fef9e7', end_color='fef9e7', fill_type='solid')
                elif ind.confirmado:
                    cell.fill = PatternFill(start_color='eafaf1', end_color='eafaf1', fill_type='solid')

            row_num += 1

    for col_num in range(1, len(headers) + 1):
        ws_consolidado.column_dimensions[get_column_letter(col_num)].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='mci_indicadores_todos.xlsx'
    )


@gestao_bp.route('/dashboard', methods=['GET'])
def dashboard():
    if not verificar_autenticacao():
        return redirect(url_for('gestao.login'))

    ciclos = CicloRevisao.query.all()
    total_ciclos = len(ciclos)
    ciclos_revistos = sum(1 for c in ciclos if c.estado == 'revisto')
    ciclos_pendentes = total_ciclos - ciclos_revistos

    ultimo_ciclo_revisto = None
    for c in sorted(ciclos, key=lambda x: x.data_revisao or datetime.min, reverse=True):
        if c.estado == 'revisto':
            ultimo_ciclo_revisto = c
            break

    confirmados_total = sum(1 for c in ciclos for i in c.indicadores if i.confirmado)
    propostas_eliminacao = sum(1 for c in ciclos for i in c.indicadores if i.propor_eliminacao)
    editados_total = sum(1 for c in ciclos for i in c.indicadores if i.foi_editado_pelo_director)

    indicadores_tipo = {'M': 0, 'A': 0}
    for c in ciclos:
        for i in c.indicadores:
            indicadores_tipo[i.tipo] += 1

    propostas_eliminacao_list = []
    for c in ciclos:
        for i in c.indicadores:
            if i.propor_eliminacao:
                propostas_eliminacao_list.append({
                    'codigo': i.codigo,
                    'nome': i.nome_original,
                    'justificacao': i.justificacao_eliminacao,
                    'ciclo': c.titulo,
                    'director': c.nome_director or '—',
                    'data': c.data_revisao.strftime('%d/%m/%Y') if c.data_revisao else '—'
                })

    edicoes_list = []
    for c in ciclos:
        for i in c.indicadores:
            if i.foi_editado_pelo_director:
                edicoes_list.append({
                    'codigo': i.codigo,
                    'nome_original': i.nome_original,
                    'nome_novo': i.nome_editado or '—',
                    'meta_original': i.meta_original,
                    'meta_nova': i.meta_editada or '—',
                    'ciclo': c.titulo,
                    'director': c.nome_director or '—',
                    'data': c.data_revisao.strftime('%d/%m/%Y') if c.data_revisao else '—'
                })

    resultados_list = []
    for c in sorted(ciclos, key=lambda x: x.criado_em, reverse=True):
        for i in c.indicadores:
            if i.resultado_actual:
                resultados_list.append({
                    'codigo': i.codigo,
                    'nome': i.nome_vigente,
                    'meta': i.meta_vigente,
                    'resultado': i.resultado_actual,
                    'ciclo': c.titulo,
                    'data': c.criado_em.strftime('%d/%m/%Y')
                })

    confirmacao_por_grupo = {}
    if ultimo_ciclo_revisto:
        for ind in ultimo_ciclo_revisto.indicadores:
            if ind.grupo not in confirmacao_por_grupo:
                confirmacao_por_grupo[ind.grupo] = {'confirmados': 0, 'total': 0}
            confirmacao_por_grupo[ind.grupo]['total'] += 1
            if ind.confirmado:
                confirmacao_por_grupo[ind.grupo]['confirmados'] += 1

    return render_template(
        'gestao/dashboard.html',
        total_ciclos=total_ciclos,
        ciclos_revistos=ciclos_revistos,
        ciclos_pendentes=ciclos_pendentes,
        confirmados_total=confirmados_total,
        propostas_eliminacao=propostas_eliminacao,
        editados_total=editados_total,
        indicadores_tipo=indicadores_tipo,
        propostas_eliminacao_list=propostas_eliminacao_list,
        edicoes_list=edicoes_list,
        resultados_list=resultados_list,
        confirmacao_por_grupo=confirmacao_por_grupo,
        ultimo_ciclo_revisto=ultimo_ciclo_revisto
    )
